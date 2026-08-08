from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "AI_Browser_Assistant_Capabilities.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MUTED = "5B6573"


def set_font(run, size=11, color="000000", bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_font(run)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_font(run)
    return p


def add_body(doc, text, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_font(run)
    return p


def add_heading(doc, text, level=1):
    style = doc.styles[f"Heading {level}"]
    p = doc.add_paragraph(style=style)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    return p


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade_cell(cell, LIGHT_GRAY)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(f"{label}  ")
    set_font(r, size=10.5, color=INK, bold=True)
    r = p.add_run(text)
    set_font(r, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for level, size, color, before, after in [
    (1, 16, BLUE, 18, 10),
    (2, 13, BLUE, 14, 7),
    (3, 12, DARK_BLUE, 10, 5),
]:
    style = doc.styles[f"Heading {level}"]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

# Quiet footer
footer_p = section.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer_p.paragraph_format.space_before = Pt(0)
footer_run = footer_p.add_run("AI Browser Assistant Capabilities")
set_font(footer_run, size=8.5, color=MUTED)

# Title block
title = doc.add_paragraph()
title.paragraph_format.space_after = Pt(3)
title.paragraph_format.keep_with_next = True
title_run = title.add_run("AI Browser Assistant")
set_font(title_run, size=24, color=INK, bold=True)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
subtitle.paragraph_format.keep_with_next = True
subtitle_run = subtitle.add_run("What it can do, where it needs your direction, and what it cannot guarantee")
set_font(subtitle_run, size=12, color=MUTED)

add_callout(
    doc,
    "Scope:",
    "This guide describes browser-based AI assistance in a shared, human-supervised workflow. Capabilities vary by website, account permissions, and browser controls.",
)

add_heading(doc, "At a glance", 1)
add_body(doc, "An AI browser assistant can help with routine web tasks, drafting, navigation, research, and form completion. It works best when you set clear rules, stay available for questions, and retain control of meaningful external actions.")

table = doc.add_table(rows=1, cols=3)
set_table_geometry(table, [2220, 3900, 3240])
headers = ["Area", "AI can help with", "Your role / limits"]
for cell, text in zip(table.rows[0].cells, headers):
    shade_cell(cell, LIGHT_BLUE)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_font(r, size=10, color=INK, bold=True)

rows = [
    ("Browser work", "Open pages, search, scroll, navigate tabs it controls, and complete multi-step workflows.", "Confirm the destination and remain available for sensitive choices or access requests."),
    ("Forms", "Enter factual information you provide, select options, and pause when required details are missing or unclear.", "Supply accurate details and decide ambiguous or consequential answers."),
    ("Writing", "Turn your notes into clear replies, posts, emails, or drafts that match a requested tone.", "Review the final wording when it represents you or your organization."),
    ("Social media", "Read public posts, summarize threads, prepare responses, and help organize engagement work.", "Approve the exact post before it is published; platform rules may also limit automation."),
    ("Longer tasks", "Work in defined steps and periodically check a page or chat while the task remains active.", "Use a scheduled monitor for repeat checks; do not assume an instant reaction to a site notification."),
]
for i, row_values in enumerate(rows):
    row = table.add_row()
    for cell, text in zip(row.cells, row_values):
        if i % 2 == 1:
            shade_cell(cell, "F8FAFC")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(text)
        set_font(r, size=9.5, color="1F2937", bold=(cell is row.cells[0]))

add_heading(doc, "What the assistant can do well", 1)
add_heading(doc, "Work through a browser task", 2)
add_bullet(doc, "Open a website, locate relevant information, search, scroll, and navigate a defined workflow.")
add_bullet(doc, "Complete routine fields using genuine information you provide, such as contact details, dates, selections, and text responses.")
add_bullet(doc, "Keep a draft form ready for your review rather than submitting it automatically.")

add_heading(doc, "Draft and refine communication", 2)
add_bullet(doc, "Convert a rough idea, voice note transcription, or bullet points into a concise message, social reply, or post.")
add_bullet(doc, "Adapt wording for a requested tone: direct, friendly, formal, brief, or explanatory.")
add_bullet(doc, "Show the final text for review before sending or posting it.")

add_heading(doc, "Ask for help at the right moment", 2)
add_body(doc, "You can set a rule in advance. For example:")
quote = doc.add_paragraph()
quote.paragraph_format.left_indent = Inches(0.25)
quote.paragraph_format.right_indent = Inches(0.25)
quote.paragraph_format.space_after = Pt(8)
quote.paragraph_format.line_spacing = 1.25
q = quote.add_run("“If a required field is missing, unclear, or needs a choice, stop and ask me. Do not guess, skip it, or submit the form.”")
set_font(q, size=10.5, color=DARK_BLUE, italic=True)
add_body(doc, "When the assistant reaches that condition, it pauses, explains what it needs, and continues after you answer.")

add_heading(doc, "Important limits", 1)
add_heading(doc, "Not truly event-driven", 2)
add_body(doc, "A browser assistant can periodically check a page or chat while a task is active. It does not reliably wake itself the instant a website notification arrives while it is idle. A scheduled monitor can improve regular check-ins, but it is still polling rather than a guaranteed push-notification response.")

add_heading(doc, "Not guaranteed to control what you see", 2)
add_body(doc, "The assistant may be able to work on a background page without changing the tab you are watching. In some browser setups, it cannot reliably activate an already-open tab in your visible shared view. This matters when you expect a video to stay visible while the assistant works elsewhere.")

add_heading(doc, "Not a substitute for your judgment", 2)
add_bullet(doc, "The assistant should not invent credentials, personal details, work history, or other factual information for real-world forms.")
add_bullet(doc, "It should pause for unclear, sensitive, high-impact, or irreversible choices instead of guessing.")
add_bullet(doc, "Website restrictions, CAPTCHA checks, logins, and access permissions can block automation.")

add_heading(doc, "Publishing and sending", 2)
add_body(doc, "Sending a message, publishing a social post, submitting a form, making a purchase, or changing access settings creates an external effect. The safe workflow is: prepare the action, show you the exact result, and obtain confirmation immediately before it happens.")

add_heading(doc, "A practical workflow", 1)
for step in [
    "Define the task and give the assistant the information it may use.",
    "Set stop rules: what to ask about, what not to guess, and which actions require approval.",
    "Let the assistant handle routine navigation, drafting, and form completion.",
    "Review drafts and answer questions as they arise.",
    "Approve the final send, post, or submission yourself.",
]:
    add_number(doc, step)

add_heading(doc, "Useful instructions to give your assistant", 1)
add_body(doc, "Clear boundaries make browser work easier to review. These are examples you can reuse or adapt:")
for label, text in [
    ("Ask-me rule:", "If a required field is missing, unclear, or has more than one reasonable answer, pause and ask me. Do not guess."),
    ("Draft-first rule:", "Prepare the response in the text box, but do not send, post, submit, purchase, upload, or change permissions until I confirm the final version."),
    ("Information rule:", "Use only the information I provide in this chat or on the approved page. Do not invent credentials, dates, qualifications, contact details, or attachments."),
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.1)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.2
    lead = p.add_run(f"{label} ")
    set_font(lead, size=10.5, color=DARK_BLUE, bold=True)
    body = p.add_run(text)
    set_font(body, size=10.5, color="1F2937")

add_callout(
    doc,
    "Best use:",
    "Treat the assistant as a fast, careful operator working alongside you—not as an unattended account owner. Clear instructions and final human review make browser assistance both more useful and safer.",
)

doc.save(OUT)
print(OUT)
